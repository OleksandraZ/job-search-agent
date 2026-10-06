"""Full job-market statistics for one or more keywords files, as an Excel workbook.

Unlike main.py's daily run this reports EVERY fitting position currently listed -
seen or not - and never sends anything or marks anything seen (storage/jobs.db is
not touched). Meant as evidence of the job market for a given profile (e.g. for the
Arbeitsagentur): one sheet per keywords file plus a summary sheet.

Fetches once with the union of all files' title_match_terms, then applies each
file's own filters (main.py:select_jobs()) to that one shared pool - four separate
fetches would repeat every company fetch four times for the same result. Does not
call resolve_pending(), so config/companies.yaml is never modified; companies still
at `ats: null` are simply skipped, as in any run before resolution.

The same job reposted on several boards (or for several cities) is merged into one
row - see pipeline/duplicates.py - so the counts are distinct positions, not listings.

Every run also saves all fetched raw jobs (descriptions included) as JSON next to
the workbook, so after a keywords change the workbook can be rebuilt from that file in
seconds with --from-raw instead of fetching everything again.

    python -m tools.market_report                     # all four keywords files
    python -m tools.market_report --keywords keywords_qa.yaml --out /tmp/qa.xlsx
    python -m tools.market_report --from-raw reports/raw_jobs_2026-10-05.json
"""

import argparse
import dataclasses
import datetime
import json
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import main
from adapters.boards import NormalizedJob
from adapters.registry import fetch_from_companies, fetch_from_sources
from agents import germany_remote, munich_local
from agents._common import SOURCE_IDS
from pipeline import classify_language
from pipeline.duplicates import duplicate_key, group_duplicates

logger = logging.getLogger(__name__)

DEFAULT_KEYWORDS_FILES = [
    "keywords_qa.yaml",
    "keywords_python.yaml",
    "keywords_cloud_devops.yaml",
    "keywords_data_engineering.yaml",
]
REPORTS_DIR = Path(__file__).resolve().parent.parent / "reports"
COLUMNS = ["Stellentitel", "Unternehmen", "Ort", "Bereich", "Sprache", "Quelle", "Link"]
# Excel sheet titles: max 31 chars, and none of these characters.
_INVALID_SHEET_CHARS = str.maketrans({c: "-" for c in "/\\?*[]:"})


def union_search_terms(keywords_configs: list[dict]) -> list[str]:
    """All files' title_match_terms, first occurrence order, case-insensitive dedupe."""
    seen: set[str] = set()
    terms: list[str] = []
    for config in keywords_configs:
        for term in config["title_match_terms"]:
            if term.lower() not in seen:
                seen.add(term.lower())
                terms.append(term)
    return terms


def _unique(values: list[str]) -> str:
    return ", ".join(dict.fromkeys(v for v in values if v))


def scope_label(job: NormalizedJob) -> str:
    munich = bool(munich_local.filter_jobs([job]))
    remote = bool(germany_remote.filter_jobs([job]))
    if munich and remote:
        return "München + Remote"
    return "München" if munich else "Remote (Deutschland)"


def language_label(job: NormalizedJob) -> str:
    return "Deutsch erforderlich" if classify_language.is_german_required(job) else "Englisch"


def sheet_title(label: str) -> str:
    return label.translate(_INVALID_SHEET_CHARS)[:31]


def build_workbook(results: list[tuple[str, list[NormalizedJob]]], as_of: datetime.date) -> Workbook:
    """results: (report_label, selected jobs) per keywords file. Pure - no I/O."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Übersicht"
    summary.append([f"Stellenmarkt-Auswertung, Stand {as_of:%d.%m.%Y}"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append(
        ["Aktuell ausgeschriebene Stellen (München und Remote in Deutschland), unabhängig "
         "vom Veröffentlichungsdatum. Dieselbe Stelle auf mehreren Portalen oder für mehrere "
         "Standorte zählt einmal."]
    )
    summary.append([])
    header = ["Suchprofil", "Stellen gesamt", "München", "Remote (Deutschland)",
              "Deutsch erforderlich", "Englisch"]
    summary.append(header)
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)

    all_keys: set[str] = set()
    for label, jobs in results:
        rows = []
        for group in group_duplicates(jobs):
            scopes = {scope_label(job) for job in group.copies}
            munich = any(s.startswith("München") for s in scopes)
            remote = any(s != "München" for s in scopes)
            if munich and remote:
                scope = "München + Remote"
            else:
                scope = "München" if munich else "Remote (Deutschland)"
            rows.append((group.main, scope, language_label(group.best_described), group.copies))
            all_keys.add(duplicate_key(group.main))

        summary.append([
            label,
            len(rows),
            sum(1 for _, scope, _, _ in rows if scope.startswith("München")),
            sum(1 for _, scope, _, _ in rows if scope != "München"),
            sum(1 for _, _, language, _ in rows if language == "Deutsch erforderlich"),
            sum(1 for _, _, language, _ in rows if language == "Englisch"),
        ])

        sheet = wb.create_sheet(sheet_title(label))
        sheet.append(COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for best, scope, language, copies in rows:
            sheet.append([
                best.title,
                best.company,
                _unique([job.location for job in copies]),
                scope,
                language,
                _unique([job.source_id for job in copies]),
                best.url,
            ])
            link = sheet.cell(row=sheet.max_row, column=len(COLUMNS))
            link.hyperlink = best.url
            link.style = "Hyperlink"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for i, width in enumerate([55, 30, 30, 22, 20, 28, 60], start=1):
            sheet.column_dimensions[get_column_letter(i)].width = width

    summary.append([])
    summary.append(["Eindeutige Stellen über alle Suchprofile", len(all_keys)])
    summary.column_dimensions["A"].width = 42
    for col in "BCDEF":
        summary.column_dimensions[col].width = 20
    return wb


def save_raw_jobs(jobs: list[NormalizedJob], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = [dataclasses.asdict(job) for job in jobs]
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def load_raw_jobs(path: Path) -> list[NormalizedJob]:
    return [NormalizedJob(**job) for job in json.loads(path.read_text(encoding="utf-8"))]


def drop_disabled_boards(jobs: list[NormalizedJob], sources_config: dict) -> list[NormalizedJob]:
    """Drop jobs from board sources no longer in SOURCE_IDS (disabled since the raw
    file was saved), so a --from-raw rebuild matches the current source list.
    Company (ATS) jobs have source ids that aren't board ids and are always kept."""
    board_ids = {source["id"] for source in sources_config["sources"]}
    return [job for job in jobs if job.source_id not in board_ids or job.source_id in SOURCE_IDS]


def fetch_raw_jobs(keywords_configs: list[dict]) -> list[NormalizedJob]:
    combined = {"title_match_terms": union_search_terms(keywords_configs)}
    logger.info("fetching once with %d combined search terms", len(combined["title_match_terms"]))
    raw_jobs = fetch_from_sources(main.load_yaml("sources.yaml"), combined, sorted(SOURCE_IDS))
    raw_jobs += fetch_from_companies(main.load_yaml("companies.yaml"), combined)
    logger.info("fetched %d raw jobs", len(raw_jobs))
    return raw_jobs


def main_cli(keywords_files: list[str], out: Path, from_raw: Path | None) -> None:
    keywords_configs = [main.load_yaml(name) for name in keywords_files]
    if from_raw:
        # Only keywords files whose title_match_terms were part of that fetch's search
        # terms can be fully re-filtered - a brand-new title term never got queried.
        raw_jobs = drop_disabled_boards(load_raw_jobs(from_raw), main.load_yaml("sources.yaml"))
        logger.info("loaded %d raw jobs from %s (no fetch)", len(raw_jobs), from_raw)
    else:
        raw_jobs = fetch_raw_jobs(keywords_configs)
        raw_path = out.with_name(out.stem.replace("market_report", "raw_jobs") + ".json")
        save_raw_jobs(raw_jobs, raw_path)
        logger.info("saved raw jobs to %s", raw_path)

    results = []
    for name, config in zip(keywords_files, keywords_configs, strict=True):
        jobs = main.select_jobs(raw_jobs, config)
        results.append((config.get("report_label", name), jobs))
        logger.info("%s: %d fitting jobs", name, len(jobs))

    out.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(results, datetime.date.today()).save(out)
    logger.info("wrote %s", out)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--keywords", nargs="+", default=DEFAULT_KEYWORDS_FILES)
    parser.add_argument(
        "--out",
        type=Path,
        default=REPORTS_DIR / f"market_report_{datetime.date.today():%Y-%m-%d}.xlsx",
    )
    parser.add_argument("--from-raw", type=Path, help="rebuild from a saved raw_jobs_*.json, no fetch")
    args = parser.parse_args()
    main_cli(args.keywords, args.out, args.from_raw)
